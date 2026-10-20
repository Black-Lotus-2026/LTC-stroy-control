"""Project boundaries, extraction and cited LangChain conclusions."""

import asyncio
import io
import json
import uuid
from datetime import UTC, datetime
from pathlib import Path

import pytest
from fastapi import HTTPException, UploadFile
from openpyxl import Workbook
from sqlalchemy import create_engine, select
from sqlalchemy.dialects.postgresql import JSONB
from sqlalchemy.ext.compiler import compiles
from sqlalchemy.orm import Session

from app.api.v1.documents import (
    AnalysisRequest,
    analyze_stage,
    delete_document,
    download_document,
    list_documents,
    upload_document,
)
from app.db.base import Base
from app.models.document import DocumentChunk, ProjectDocument
from app.models.enums import UserRole
from app.models.project import Project
from app.models.schedule import ScheduleTask
from app.models.user import User
from app.services.document_search import (
    AgentUnavailable,
    MachineryAnswer,
    build_chain,
    extract_chunks,
    search_machinery,
    select_evidence,
)
from app.services.document_workflow import DocumentMachineryWorkflow


@compiles(JSONB, "sqlite")
def compile_jsonb_sqlite(type_, compiler, **kw):
    return "JSON"


@pytest.fixture
def db():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    with Session(engine) as session:
        yield session
    engine.dispose()


def setup(db):
    owner = User(username="docs-owner", name="Владелец", role=UserRole.MANAGER)
    outsider = User(username="docs-outsider", name="Чужой", role=UserRole.MANAGER)
    project = Project(code="DOCS-ONE", name="Первый", owner_username=owner.username)
    other = Project(code="DOCS-TWO", name="Второй", owner_username=outsider.username)
    stage = ScheduleTask(
        project=project,
        name="Разработка грунта",
        order_index=1,
        planned_start=datetime.now(UTC),
        planned_end=datetime.now(UTC),
    )
    db.add_all([owner, outsider, project, other, stage])
    db.commit()
    return owner, outsider, project, other, stage


def test_extract_docx_with_locator():
    pytest.importorskip("docx")
    from docx import Document as WordDocument

    word = WordDocument()
    word.add_paragraph("Разработка грунта: требуется экскаватор и самосвал")
    stream = io.BytesIO()
    word.save(stream)
    chunks = extract_chunks(stream.getvalue(), ".docx")
    assert chunks[0][0] == "абзац 1"
    assert "экскаватор" in chunks[0][1]


def test_extract_xlsx_with_locator():
    book = Workbook()
    book.active.title = "Расценки"
    book.active.append(["Земляные работы", "Экскаватор", "2 маш.-ч"])
    stream = io.BytesIO()
    book.save(stream)
    chunks = extract_chunks(stream.getvalue(), ".xlsx")
    assert chunks[0][0] == "лист Расценки, строка 1"
    assert "Экскаватор" in chunks[0][1]


def test_extract_pdf_pages():
    pytest.importorskip("pypdfium2")
    sample = Path(__file__).parents[2] / "docs" / "technical_documentation.pdf"
    chunks = extract_chunks(sample.read_bytes(), ".pdf")
    assert chunks[0][0] == "стр. 1"
    assert len(chunks[0][1]) > 100


def test_stage_heading_includes_next_row(db):
    _, _, _, _, stage = setup(db)
    document_id = uuid.uuid4()
    chunks = [
        DocumentChunk(
            id=uuid.uuid4(),
            document_id=document_id,
            ordinal=0,
            locator="лист 1, строка 1",
            text="Разработка грунта",
        ),
        DocumentChunk(
            id=uuid.uuid4(),
            document_id=document_id,
            ordinal=1,
            locator="лист 1, строка 2",
            text="Экскаватор 2 машино-часа",
        ),
    ]
    assert select_evidence(stage, chunks) == chunks


def test_retrieval_covers_multiple_documents(db):
    _, _, _, _, stage = setup(db)
    first = uuid.uuid4()
    second = uuid.uuid4()
    chunks = [
        DocumentChunk(
            id=uuid.uuid4(),
            document_id=first,
            ordinal=index,
            locator=f"стр. {index + 1}",
            text="Разработка грунта экскаватором",
        )
        for index in range(10)
    ]
    chunks.append(
        DocumentChunk(
            id=uuid.uuid4(),
            document_id=second,
            ordinal=0,
            locator="стр. 1",
            text="Разработка грунта самосвалами",
        )
    )
    chosen = select_evidence(stage, chunks, limit=5)
    assert {chunk.document_id for chunk in chosen} == {first, second}


def test_agent_plans_synonym_search_when_stage_words_do_not_match(db):
    _, _, _, _, stage = setup(db)
    chunk = DocumentChunk(
        id=uuid.uuid4(),
        document_id=uuid.uuid4(),
        ordinal=0,
        locator="стр. 7",
        text="Земляные работы: требуется экскаватор",
    )

    class Planner:
        async def ainvoke(self, value):
            assert value["stage"] == stage.name
            return {"search_terms": ["земляные работы"]}

    chain = FakeChain(
        {
            "items": [
                {
                    "name": "Экскаватор",
                    "purpose": "",
                    "chunk_id": str(chunk.id),
                    "quote": "требуется экскаватор",
                    "requirement": "required",
                    "suggested_code": "MACHINERY_EXCAVATOR",
                }
            ],
            "summary": "Найдено",
        }
    )
    result = asyncio.run(search_machinery(stage, [chunk], chain=chain, planner=Planner()))
    assert result.items[0]["locator"] == "стр. 7"


def test_upload_download_and_project_isolation(db):
    owner, outsider, project, other, stage = setup(db)
    data = "Разработка грунта: необходим экскаватор".encode()
    upload = UploadFile(filename="ППР.txt", file=io.BytesIO(data))
    info = asyncio.run(upload_document(project.id, db, owner, upload, "PPR"))
    assert info.chunk_count == 1
    assert list_documents(project.id, db, owner)[0].filename == "ППР.txt"
    assert download_document(project.id, info.id, db, owner).body == data
    with pytest.raises(HTTPException) as denied:
        list_documents(project.id, db, outsider)
    assert denied.value.status_code == 404
    with pytest.raises(HTTPException) as denied:
        download_document(other.id, info.id, db, outsider)
    assert denied.value.status_code == 404
    with pytest.raises(HTTPException) as denied:
        asyncio.run(
            analyze_stage(project.id, AnalysisRequest(stage_id=uuid.uuid4()), db, owner)
        )
    assert denied.value.status_code == 404
    foreign_stage = ScheduleTask(
        project=other,
        name="Монтаж конструкций",
        order_index=1,
        planned_start=datetime.now(UTC),
        planned_end=datetime.now(UTC),
    )
    db.add(foreign_stage)
    db.commit()
    with pytest.raises(HTTPException) as denied:
        asyncio.run(
            analyze_stage(
                project.id, AnalysisRequest(stage_id=foreign_stage.id), db, owner
            )
        )
    assert denied.value.status_code == 404
    with pytest.raises(HTTPException) as duplicate:
        asyncio.run(
            upload_document(
                project.id,
                db,
                owner,
                UploadFile(filename="copy.txt", file=io.BytesIO(data)),
                "PPR",
            )
        )
    assert duplicate.value.status_code == 409
    delete_document(project.id, info.id, db, owner)
    assert db.scalars(select(DocumentChunk)).all() == []


class FakeChain:
    def __init__(self, answer):
        self.answer = answer
        self.input = None

    async def ainvoke(self, value):
        self.input = value
        return self.answer


def test_workflow_scopes_sources_and_releases_db_transaction(db):
    owner, outsider, project, _, stage = setup(db)
    doc = ProjectDocument(
        project_id=project.id,
        kind="POS",
        filename="ПОС.txt",
        content_type="text/plain",
        size_bytes=50,
        sha256="b" * 64,
        content=b"source",
    )
    doc.chunks = [
        DocumentChunk(
            ordinal=0,
            locator="стр. 2",
            text="Разработка грунта: нужен экскаватор",
        )
    ]
    db.add(doc)
    db.commit()
    chunk_id = doc.chunks[0].id
    chain = FakeChain(
        MachineryAnswer.model_validate(
            {
                "items": [
                    {
                        "name": "Экскаватор",
                        "purpose": "разработка грунта",
                        "chunk_id": str(chunk_id),
                        "quote": "нужен экскаватор",
                    }
                ],
                "summary": "Найдена техника",
            }
        )
    )
    workflow = DocumentMachineryWorkflow(db, chain=chain)
    result = asyncio.run(
        workflow.run(project_id=project.id, stage_id=stage.id, user=owner)
    )
    assert result["items"][0]["document_name"] == "ПОС.txt"
    assert result["items"][0]["locator"] == "стр. 2"
    assert db.in_transaction() is False
    with pytest.raises(HTTPException) as denied:
        asyncio.run(workflow.run(project_id=project.id, stage_id=stage.id, user=outsider))
    assert denied.value.status_code == 404


def test_agent_accepts_only_exact_cited_machinery(db):
    owner, _, project, _, stage = setup(db)
    doc = ProjectDocument(
        project_id=project.id,
        kind="PPR",
        filename="ППР.txt",
        content_type="text/plain",
        size_bytes=50,
        sha256="a" * 64,
        content=b"source",
    )
    doc.chunks = [
        DocumentChunk(
            ordinal=0,
            locator="стр. 5",
            text="Разработка грунта: необходим экскаватор 2 шт.",
        )
    ]
    db.add(doc)
    db.commit()
    chunk = doc.chunks[0]
    chain = FakeChain(
        MachineryAnswer.model_validate(
            {
                "items": [
                    {
                        "name": "Экскаватор",
                        "purpose": "разработка грунта",
                        "chunk_id": str(chunk.id),
                        "quote": "необходим экскаватор 2 шт.",
                    },
                    {
                        "name": "Автокран",
                        "purpose": "подъём",
                        "chunk_id": str(chunk.id),
                        "quote": "необходим автокран",
                    },
                    {
                        "name": "Самосвал",
                        "purpose": "вывоз",
                        "chunk_id": str(uuid.uuid4()),
                        "quote": "самосвал",
                    },
                ],
                "summary": "Найдена техника",
            }
        )
    )
    result = asyncio.run(search_machinery(stage, doc.chunks, chain=chain))
    assert len(result.items) == 1
    assert result.items[0]["name"] == "Экскаватор"
    assert result.items[0]["locator"] == "стр. 5"
    assert "Разработка грунта" in chain.input["facts_json"]


def test_agent_builds_external_api_chain(monkeypatch):
    pytest.importorskip("langchain_openai")
    from pydantic import SecretStr

    monkeypatch.setattr(
        "app.services.document_search.settings.document_llm_base_url",
        "https://example-model-provider.test/v1",
    )
    monkeypatch.setattr(
        "app.services.document_search.settings.document_llm_model", "remote-test-model"
    )
    monkeypatch.setattr(
        "app.services.document_search.settings.document_llm_api_key",
        SecretStr("test-key"),
    )
    chain = build_chain()
    assert chain.__class__.__name__ == "RunnableSequence"
    assert str(chain.steps[1].openai_api_base) == "https://example-model-provider.test/v1"
    assert chain.steps[1].model_name == "remote-test-model"


def test_agent_requires_api_key(monkeypatch):
    from pydantic import SecretStr

    monkeypatch.setattr(
        "app.services.document_search.settings.document_llm_api_key", SecretStr("")
    )
    monkeypatch.setattr(
        "app.services.document_search.settings.openrouter_api_key", SecretStr("")
    )
    with pytest.raises(AgentUnavailable):
        build_chain()


def test_agent_rejects_local_api_endpoint(monkeypatch):
    from pydantic import SecretStr

    monkeypatch.setattr(
        "app.services.document_search.settings.document_llm_api_key",
        SecretStr("test-key"),
    )
    monkeypatch.setattr(
        "app.services.document_search.settings.document_llm_base_url",
        "http://127.0.0.1:11434/v1",
    )
    with pytest.raises(AgentUnavailable):
        build_chain()


def test_langchain_parses_local_chat_response(db):
    pytest.importorskip("langchain_core")
    from langchain_core.language_models.fake_chat_models import FakeListChatModel

    _, _, _, _, stage = setup(db)
    document_id = uuid.uuid4()
    chunk = DocumentChunk(
        id=uuid.uuid4(),
        document_id=document_id,
        ordinal=0,
        locator="стр. 3",
        text="Разработка грунта: требуется экскаватор",
    )
    llm = FakeListChatModel(
        responses=[
            json.dumps(
                {
                    "items": [
                        {
                            "name": "Экскаватор",
                            "purpose": "разработка грунта",
                            "chunk_id": str(chunk.id),
                            "quote": "требуется экскаватор",
                        }
                    ],
                    "summary": "Техника найдена",
                },
                ensure_ascii=False,
            )
        ]
    )
    result = asyncio.run(search_machinery(stage, [chunk], chain=build_chain(llm=llm)))
    assert result.items[0]["document_id"] == document_id
    assert result.summary == "Техника найдена"
