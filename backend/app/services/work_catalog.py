"""Импорт справочника видов работ из файла организаторов.

Файл «Сводный перечень строительных работ» — это иерархический справочник
на 378 наименований, а не календарный план: ни сроков, ни зон в нём нет.
Он даёт ответ на вопрос «какие работы бывают», на который потом опираются
правила «этап работ → необходимая техника».

Разбор осложняют две особенности исходного файла, обе учтены ниже:

1. Excel превратил часть кодов в даты. Код «10.1.» сохранён в файле как
   число 45667 — это 10 января 2025 года. Восстанавливаем обратной
   подстановкой «день.месяц».
2. Уровень вложенности размечен не полностью: код заполнен у 128 строк из
   378, а стандартной группировки строк (outline level) в файле нет.
   Строки без кода — это подпункты предыдущей строки с кодом.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import IO

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import WorkType

# Строка с заголовками колонок. Выше неё — название справочника.
HEADER_ROW = 3
# Колонка с номером пункта и колонка с наименованием работы.
CODE_COLUMN = 1
NAME_COLUMN = 2
# Начиная с этой колонки идут типы объектов («Жильё», «Образование», ...),
# отмеченные галочками.
FIRST_APPLICABILITY_COLUMN = 3

# Символ, которым в файле отмечена применимость вида работ к типу объекта.
APPLICABILITY_MARKS = {"˅", "v", "V", "✓", "+", "х", "x", "X"}

LEVEL_BY_CODE_PATTERN = (
    (re.compile(r"^\d+\.?$"), 1),
    (re.compile(r"^\d+\.\d+\.?$"), 2),
    (re.compile(r"^\d+\.\d+\.\d+\.?$"), 3),
    (re.compile(r"^\d+\.\d+\.\d+\.\d+\.?$"), 4),
)


@dataclass
class ParsedWorkType:
    """Одна строка справочника после разбора."""

    name: str
    level: int
    source_row: int
    code: str | None = None
    # Индекс родителя в том же списке; None — корневой раздел.
    parent_index: int | None = None
    applicable_to: list[str] = field(default_factory=list)
    # Код восстановлен из даты, в которую его превратил Excel.
    code_recovered_from_date: bool = False


@dataclass
class ParseReport:
    """Итог разбора файла — показывается пользователю после импорта."""

    items: list[ParsedWorkType]
    recovered_codes: int = 0
    rows_without_code: int = 0
    skipped_rows: int = 0

    @property
    def total(self) -> int:
        return len(self.items)

    @property
    def by_level(self) -> dict[int, int]:
        counts: dict[int, int] = {}
        for item in self.items:
            counts[item.level] = counts.get(item.level, 0) + 1
        return dict(sorted(counts.items()))


def _restore_code_from_date(value: datetime | date) -> str:
    """Вернуть код пункта, который Excel сохранил как дату.

    «10.1.» при открытии файла воспринимается как 10 января и хранится
    числом. Обратное преобразование — «день.месяц».
    """
    return f"{value.day}.{value.month}."


def _normalize_code(raw: object) -> tuple[str | None, bool]:
    """Привести значение колонки с номером пункта к коду.

    Возвращает сам код и признак того, что он был восстановлен из даты.
    """
    if raw is None:
        return None, False
    if isinstance(raw, datetime | date):
        return _restore_code_from_date(raw), True

    text = str(raw).strip()
    if not text:
        return None, False
    # Числа без дробной части Excel отдаёт как «12.0» — приводим к «12».
    if re.fullmatch(r"\d+\.0", text):
        text = text[:-2]
    return text, False


def _level_from_code(code: str) -> int | None:
    for pattern, level in LEVEL_BY_CODE_PATTERN:
        if pattern.match(code):
            return level
    return None


def _read_applicability_headers(sheet) -> dict[int, str]:
    """Прочитать названия типов объектов из шапки таблицы."""
    headers: dict[int, str] = {}
    for column in range(FIRST_APPLICABILITY_COLUMN, sheet.max_column + 1):
        value = sheet.cell(row=HEADER_ROW, column=column).value
        if value and str(value).strip():
            headers[column] = str(value).strip()
    return headers


def parse_work_catalog(source: str | Path | IO[bytes]) -> ParseReport:
    """Разобрать xlsx-справочник в плоский список с восстановленной иерархией.

    Функция не обращается к базе: её можно запустить на любом файле и
    проверить результат отдельно от импорта.
    """
    workbook = load_workbook(source, data_only=True, read_only=True)
    sheet = workbook[workbook.sheetnames[0]]

    applicability_headers = _read_applicability_headers(sheet)

    report = ParseReport(items=[])
    # Последняя встреченная строка с кодом на каждом уровне — по ней
    # определяется родитель следующих строк.
    last_index_by_level: dict[int, int] = {}
    last_coded_level = 1

    for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if row_number <= HEADER_ROW:
            continue

        name_value = row[NAME_COLUMN - 1] if len(row) >= NAME_COLUMN else None
        name = str(name_value).strip() if name_value else ""
        if not name:
            report.skipped_rows += 1
            continue

        raw_code = row[CODE_COLUMN - 1] if len(row) >= CODE_COLUMN else None
        code, recovered = _normalize_code(raw_code)

        if code:
            level = _level_from_code(code)
            if level is None:
                # Код есть, но формат незнакомый: считаем строку подпунктом
                # предыдущего уровня, не теряя её.
                level = last_coded_level + 1
            last_coded_level = level
            if recovered:
                report.recovered_codes += 1
        else:
            # Строка без номера — подпункт последней пронумерованной строки.
            level = last_coded_level + 1
            report.rows_without_code += 1

        applicable_to = [
            header
            for column, header in applicability_headers.items()
            if len(row) >= column
            and str(row[column - 1]).strip() in APPLICABILITY_MARKS
        ]

        item = ParsedWorkType(
            name=name,
            level=level,
            source_row=row_number,
            code=code,
            parent_index=last_index_by_level.get(level - 1),
            applicable_to=applicable_to,
            code_recovered_from_date=recovered,
        )
        report.items.append(item)

        index = len(report.items) - 1
        last_index_by_level[level] = index
        # Уровни глубже текущего больше не могут быть родителями.
        for deeper in [lvl for lvl in last_index_by_level if lvl > level]:
            del last_index_by_level[deeper]

    workbook.close()
    return report


@dataclass
class ImportStats:
    """Итог записи справочника в базу."""

    created: int = 0
    updated: int = 0
    deleted: int = 0
    recovered_codes: int = 0
    rows_without_code: int = 0
    by_level: dict[int, int] = field(default_factory=dict)


def import_work_catalog(
    db: Session,
    source: str | Path | IO[bytes],
    *,
    replace: bool = False,
) -> ImportStats:
    """Разобрать файл и записать справочник в базу.

    Повторный импорт того же файла не создаёт дублей: строки сопоставляются
    по номеру строки исходного файла, который хранится в `source_row`.

    `replace=True` удаляет виды работ, которых в файле больше нет. Операция
    каскадно удалит и правила, ссылающиеся на эти виды работ, поэтому по
    умолчанию выключена.
    """
    report = parse_work_catalog(source)

    existing: dict[int, WorkType] = {
        work_type.source_row: work_type
        for work_type in db.scalars(select(WorkType)).all()
        if work_type.source_row is not None
    }

    stats = ImportStats(
        recovered_codes=report.recovered_codes,
        rows_without_code=report.rows_without_code,
        by_level=report.by_level,
    )

    # Первый проход: создаём или обновляем записи без связей с родителями.
    # Родители проставляются вторым проходом, когда у всех строк уже есть id.
    models: list[WorkType] = []
    for item in report.items:
        work_type = existing.get(item.source_row)
        if work_type is None:
            work_type = WorkType(source_row=item.source_row)
            db.add(work_type)
            stats.created += 1
        else:
            stats.updated += 1

        work_type.code = item.code
        work_type.name = item.name
        work_type.level = item.level
        work_type.sort_order = item.source_row
        work_type.applicable_to = item.applicable_to
        models.append(work_type)

    # id нужны до проставления родителей.
    db.flush()

    for item, work_type in zip(report.items, models, strict=True):
        parent = models[item.parent_index] if item.parent_index is not None else None
        work_type.parent_id = parent.id if parent else None

    if replace:
        imported_rows = {item.source_row for item in report.items}
        for source_row, work_type in existing.items():
            if source_row not in imported_rows:
                db.delete(work_type)
                stats.deleted += 1

    db.flush()
    return stats
